#!/usr/bin/env python3
"""
extract_complex.py

Extracts PO line-item data from numbered Purchase Order Excel files and
appends them to a user-selected Consolidated database workbook.
"""

import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

SOURCE_FOLDER_NAME = "Extract Complex POs in Consolidated"
CONSOLIDATED_FOLDER_HINT = "consolidat"  # matched case-insensitively against folder names

# Labels as they appear on the source PO sheets -> internal field name
HEADER_LABELS = {
    "description": "description",
    "quantity": "quantity",
    "uom": "uom",
    "unit cost": "unit_cost",
    "delivery date": "delivery_date",
    "line": "line",
}

# Labels as they appear on the Consolidated database sheets -> internal field name.
DB_HEADER_LABELS = {
    "po no.": "po_number",
    "line": "line",
    "description": "description",
    "qty": "quantity",
    "uom": "uom",
    "unit cost": "unit_cost",
    "edd": "delivery_date",
    "extended cost": "extended_cost",
    "company": "company",
    "client": "client",
}
DB_REQUIRED = {"po_number", "line", "description", "quantity", "uom", "unit_cost"}


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------

def normalize(text):
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


# ---------------------------------------------------------------------------
# reading a source PO workbook
# ---------------------------------------------------------------------------

def find_header_columns(ws):
    """Scan the sheet for the header labels and record which column each lives in."""
    found = {}
    target_keys = set(HEADER_LABELS.values())
    for row in ws.iter_rows():
        for cell in row:
            key = HEADER_LABELS.get(normalize(cell.value))
            if key and key not in found:
                found[key] = cell.column
        if target_keys.issubset(found.keys()):
            break
    return found


def is_header_row(ws, row_idx, cols):
    """True if this row is one of the repeating page-header rows (not a data row)."""
    for col in cols.values():
        if normalize(ws.cell(row=row_idx, column=col).value) in HEADER_LABELS:
            return True
    return False


def find_furniture_rows(ws):
    max_row = ws.max_row
    max_col = ws.max_column

    header_rows = []
    marker_rows = []
    for r in range(1, max_row + 1):
        is_header = False
        is_marker = False
        for c in range(1, max_col + 1):
            text = normalize(ws.cell(row=r, column=c).value)
            if not text:
                continue
            if text in HEADER_LABELS:
                is_header = True
            if text.startswith("billing address"):
                is_marker = True
        if is_header:
            header_rows.append(r)
        if is_marker:
            marker_rows.append(r)

    furniture = set()
    for start in marker_rows:
        following = [h for h in header_rows if h >= start]
        if not following:
            continue
        end = following[0]
        idx = header_rows.index(end)
        if idx + 1 < len(header_rows) and header_rows[idx + 1] == end + 1:
            end = header_rows[idx + 1]
        furniture.update(range(start, end + 1))

    return furniture


def extract_records(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    records = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        cols = find_header_columns(ws)
        required = {"quantity", "unit_cost", "uom", "description"}
        if not required.issubset(cols):
            continue

        desc_col = cols["description"]
        qty_col = cols["quantity"]
        uom_col = cols["uom"]
        cost_col = cols["unit_cost"]
        date_col = cols.get("delivery_date")
        line_col = cols.get("line")

        furniture = find_furniture_rows(ws)

        def is_data_row(row_idx):
            q = ws.cell(row=row_idx, column=qty_col).value
            u = ws.cell(row=row_idx, column=uom_col).value
            c = ws.cell(row=row_idx, column=cost_col).value
            return is_number(q) and bool(u) and is_number(c)

        for r in range(2, ws.max_row + 1):
            if r in furniture or is_header_row(ws, r, cols):
                continue
            if not is_data_row(r):
                continue

            qty_val = ws.cell(row=r, column=qty_col).value
            uom_val = ws.cell(row=r, column=uom_col).value
            cost_val = ws.cell(row=r, column=cost_col).value

            desc_parts = []
            line_no = None
            delivery_date = None
            cur = r - 1
            steps = 0
            while cur >= 1 and steps < 60:
                steps += 1
                if cur in furniture:
                    cur -= 1
                    continue
                if is_data_row(cur):
                    break

                desc_val = ws.cell(row=cur, column=desc_col).value
                if desc_val not in (None, ""):
                    desc_parts.insert(0, str(desc_val))

                if line_col:
                    l_val = ws.cell(row=cur, column=line_col).value
                    if l_val not in (None, ""):
                        line_no = l_val

                if date_col:
                    d_val = ws.cell(row=cur, column=date_col).value
                    if d_val not in (None, ""):
                        delivery_date = d_val

                if line_no is not None:
                    break
                cur -= 1

            description = " ".join(part.strip() for part in desc_parts).strip()
            if not description:
                continue

            records.append({
                "line": line_no,
                "description": description,
                "quantity": qty_val,
                "uom": str(uom_val).strip(),
                "unit_cost": cost_val,
                "delivery_date": delivery_date,
            })

    return records


def extract_po_number(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    raw_value = None

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if normalize(cell.value) != "purchase order number":
                    continue
                col, r = cell.column, cell.row
                for rr in range(r + 1, r + 4):
                    v = ws.cell(row=rr, column=col).value
                    if v not in (None, ""):
                        raw_value = str(v)
                        break
                if raw_value:
                    break
            if raw_value:
                break
        if raw_value:
            break

    if not raw_value:
        m = re.search(r"(\d{4,})", path.stem)
        if m:
            raw_value = m.group(1)

    if not raw_value:
        return None

    digits = re.findall(r"\d+", raw_value)
    return int(digits[-1]) if digits else None


def extract_supplier(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if normalize(cell.value) == "supplier":
                    below = ws.cell(row=cell.row + 1, column=cell.column).value
                    if below not in (None, ""):
                        return str(below)
    return None


# ---------------------------------------------------------------------------
# locating folders & interactive prompt UI
# ---------------------------------------------------------------------------

def find_named_folder(start_dir, folder_name):
    start_dir = start_dir.resolve()
    target = folder_name.lower()

    direct = start_dir / folder_name
    if direct.is_dir():
        return direct
    for child in start_dir.iterdir():
        if child.is_dir() and child.name.lower() == target:
            return child

    for p in start_dir.rglob("*"):
        if p.is_dir() and p.name.lower() == target:
            return p

    for parent in start_dir.parents:
        for child in parent.iterdir():
            if child.is_dir() and child.name.lower() == target:
                return child
        for p in parent.rglob("*"):
            if p.is_dir() and p.name.lower() == target:
                return p

    return None


def find_consolidated_folder(start_dir, source_folder):
    search_roots = [start_dir.resolve()]
    if source_folder is not None:
        search_roots.append(source_folder.parent)

    for root in search_roots:
        try:
            children = list(root.iterdir())
        except OSError:
            continue
        for child in children:
            if child.is_dir() and CONSOLIDATED_FOLDER_HINT in child.name.lower():
                return child
    return None


def find_database_files(consolidated_folder):
    return sorted(
        p for p in consolidated_folder.rglob("*.xlsx")
        if not p.name.startswith("~$") and CONSOLIDATED_FOLDER_HINT in p.stem.lower()
    )


def find_source_files(folder):
    return sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith("~$"))


def prompt_select_database(db_files):
    """Presents a numbered menu matching your specified layout format."""
    print("Select Company (Type the corresponding number PRESS ENTER PRESS START PRESS ENTER)\n")
    print("Select target database for ALL files:")
    for idx, db_path in enumerate(db_files, start=1):
        print(f"  {idx}. {db_path.name}")

    while True:
        choice = input(f"Enter choice (1-{len(db_files)}): ").strip()
        if choice.isdigit():
            idx = int(choice)
            if 1 <= idx <= len(db_files):
                return db_files[idx - 1]
        print(f"Invalid selection. Enter a number between 1 and {len(db_files)}.")


# ---------------------------------------------------------------------------
# reading / writing a Consolidated database
# ---------------------------------------------------------------------------

def find_db_header(ws, max_scan_rows=10):
    for r in range(1, max_scan_rows + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            key = DB_HEADER_LABELS.get(normalize(ws.cell(row=r, column=c).value))
            if key and key not in found:
                found[key] = c
        if DB_REQUIRED.issubset(found.keys()):
            return r, found
    return None, {}


def find_last_data_row(ws, header_row, anchor_col):
    for r in range(ws.max_row, header_row - 1, -1):
        if ws.cell(row=r, column=anchor_col).value not in (None, ""):
            return r
    return header_row


def default_company_client(ws, header_row, cols):
    company = client = None
    if "company" in cols or "client" in cols:
        for r in range(header_row + 1, min(ws.max_row, header_row + 20) + 1):
            if "company" in cols and company is None:
                v = ws.cell(row=r, column=cols["company"]).value
                if v not in (None, ""):
                    company = v
            if "client" in cols and client is None:
                v = ws.cell(row=r, column=cols["client"]).value
                if v not in (None, ""):
                    client = v
    return company, client


def append_entries_to_database(db_path, entries, supplier_text):
    wb = openpyxl.load_workbook(db_path, data_only=False)
    ws = wb.active

    header_row, cols = find_db_header(ws)
    if header_row is None:
        print(f"    Could not find a recognizable header row in {db_path.name}, skipping.")
        return None

    anchor_col = cols["line"]
    last_row = find_last_data_row(ws, header_row, anchor_col)
    company, client = default_company_client(ws, header_row, cols)
    if company is None and supplier_text:
        company = supplier_text.splitlines()[0].strip()

    for entry in entries:
        new_row = last_row + 1
        if "company" in cols and company is not None:
            ws.cell(row=new_row, column=cols["company"], value=company)
        if "client" in cols and client is not None:
            ws.cell(row=new_row, column=cols["client"], value=client)
        ws.cell(row=new_row, column=cols["po_number"], value=entry["po_number"])
        ws.cell(row=new_row, column=cols["line"], value=entry["line"])
        ws.cell(row=new_row, column=cols["description"], value=entry["description"])
        ws.cell(row=new_row, column=cols["quantity"], value=entry["quantity"])
        ws.cell(row=new_row, column=cols["uom"], value=entry["uom"])
        ws.cell(row=new_row, column=cols["unit_cost"], value=entry["unit_cost"])
        if "delivery_date" in cols and entry["delivery_date"] is not None:
            ws.cell(row=new_row, column=cols["delivery_date"], value=entry["delivery_date"])
        if "extended_cost" in cols and is_number(entry["quantity"]) and is_number(entry["unit_cost"]):
            ws.cell(
                row=new_row, column=cols["extended_cost"],
                value=round(entry["quantity"] * entry["unit_cost"], 2),
            )
        last_row = new_row

    wb.save(db_path)
    return last_row


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def run_once(script_dir):
    source_folder = find_named_folder(script_dir, SOURCE_FOLDER_NAME)
    if source_folder is None:
        print(f"\nCould not find a folder named '{SOURCE_FOLDER_NAME}'.")
        print("Make sure it exists near this script.")
        return

    source_files = find_source_files(source_folder)
    if not source_files:
        print(f"No .xlsx files found in '{source_folder}'.")
        return

    consolidated_folder = find_consolidated_folder(script_dir, source_folder)
    if consolidated_folder is None:
        print("\nCould not find the Consolidated folder.")
        return

    db_files = find_database_files(consolidated_folder)
    if not db_files:
        print(f"No Consolidated_*.xlsx database files found in '{consolidated_folder}'.")
        return

    # Direct selection of destination database
    target_db = prompt_select_database(db_files)

    pending = {}

    print(f"\nProcessing {len(source_files)} source file(s) in '{source_folder.name}':")
    for f in source_files:
        po_number = extract_po_number(f)
        if po_number is None:
            print(f"  {f.name}: could not determine a PO number, skipping")
            continue

        supplier_text = extract_supplier(f)
        records = extract_records(f)
        if not records:
            print(f"  {f.name}: no line items found, skipping")
            continue

        for rec in records:
            rec["po_number"] = po_number
        pending.setdefault(target_db, []).append((f.name, po_number, records, supplier_text))
        print(f"  {f.name}: PO #{po_number} -> {target_db.name}, {len(records)} line item(s)")

    if not pending:
        print("\nNo line items to add.")
        return

    print()
    for db_path, items in pending.items():
        all_entries = []
        supplier_text = None
        for fname, po_number, records, sup in items:
            all_entries.extend(records)
            supplier_text = supplier_text or sup
        last_row = append_entries_to_database(db_path, all_entries, supplier_text)
        if last_row is not None:
            print(
                f"Appended {len(all_entries)} row(s) from {len(items)} PO(s) to "
                f"'{db_path.name}' (now ends at row {last_row})."
            )


def get_app_dir():
    """
    Folder where the EXE is located when frozen with PyInstaller.
    For normal Python execution, use the folder containing this .py file.
    """
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def main():
    # IMPORTANT: with a PyInstaller EXE, __file__ can point to PyInstaller's
    # temporary extraction folder. Use the actual EXE folder instead.
    script_dir = get_app_dir()

    while True:
        run_once(script_dir)
        choice = input("Type EXIT  then PRESS  enter to quit, \n").strip().lower()
        if choice == "exit":
            print("Exitting")
            break


if __name__ == "__main__":
    main()