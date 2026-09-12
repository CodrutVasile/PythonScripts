"""
Delivery Note generator
========================
Reads item lines for a PO number out of the "Consolidated" database file and
writes them into a copy of a supplier's Delivery Note template, using Excel
itself (via COM) so that every bit of original formatting, merged cells and
embedded images survive untouched.

Works with any of the 4 current templates (Abraj, Ameedat, Athab, Wasmat)
because they all share the same table layout:

    Row 14      supplier "PO#" label            (e.g. "ABRAJ AL-DAR PO#")
    Row 15      column titles (SR#, PO#, ...)
    Row 16      first (blank, pre-formatted) item row
    Row 17+     rest of the template (signature block / footer)

To add item rows beyond the first one, we insert new rows at 17 (pushing the
footer down) and copy row 16's formatting onto them.

Run it, pick a template, pick the consolidated file, then type one or more
PO numbers.
"""

from pathlib import Path
import re
import sys
import time
import zipfile

import pandas as pd
import openpyxl
import win32com.client as win32

# ---------------------------------------------------------------------------
# Where things live
# ---------------------------------------------------------------------------

if getattr(sys, "frozen", False):
    BASE = Path(sys.executable).resolve().parent
else:
    BASE = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# Template layout (same for all 4 supplier templates)
# ---------------------------------------------------------------------------

HEADER_ROW = 15         # column titles (SR#, PO#, DESCRIPTION, ...)
ITEM_ROW = 16            # first data row
FOOTER_ROW = 17          # row where extra item rows get inserted

TABLE_FIRST_COL = 2      # column B
TABLE_LAST_COL = 12      # column L

PO_LABEL_SEARCH_ROWS = 14   # search rows 1..14 for "PO NUMBER" / "PO#" labels
PO_LABEL_PATTERN = re.compile(r"PO\s*(NUMBER|#)", re.IGNORECASE)


# ---------------------------------------------------------------------------
# File discovery (template + consolidated/database file)
# ---------------------------------------------------------------------------

def find_xlsx_files(root: Path) -> list[Path]:
    """Recursively collect all .xlsx files under root (skips Excel temp files)."""
    files = [p for p in root.rglob("*.xlsx") if not p.name.startswith("~$")]
    return sorted(files, key=lambda p: str(p).lower())


def is_template(path: Path) -> bool:
    """Heuristic: filename looks like a Delivery Note template."""
    name = path.name.upper()
    return ("TEMPLATE" in name and "DELIVERY" in name) or (
        "TEMPLATE" in name and "NOTE" in name
    )


def is_consolidated(path: Path) -> bool:
    """Strict: filename must contain 'Consolidated'."""
    return "CONSOLIDATED" in path.name.upper()


def relative_display(path: Path) -> str:
    """Show path relative to BASE for cleaner menus."""
    try:
        return str(path.relative_to(BASE))
    except ValueError:
        return str(path)


def choose_from_list(title: str, items: list[Path]) -> Path:
    """Present a numbered menu (grouped by folder) and return the chosen Path."""
    if not items:
        raise RuntimeError(f"No files found for: {title}")

    print()
    print("=" * 70)
    print(title)
    print("=" * 70)

    by_folder: dict[str, list[Path]] = {}
    for p in items:
        folder = relative_display(p.parent) if p.parent != BASE else "."
        by_folder.setdefault(folder, []).append(p)

    flat: list[Path] = []
    idx = 1
    for folder in sorted(by_folder.keys(), key=str.lower):
        print(f"\n  [{folder}]")
        for p in by_folder[folder]:
            print(f"    {idx:4d}.  {p.name}")
            flat.append(p)
            idx += 1

    print()
    while True:
        raw = input(f"Enter number (1-{len(flat)}): ").strip()
        if not raw:
            continue
        try:
            choice = int(raw)
            if 1 <= choice <= len(flat):
                selected = flat[choice - 1]
                print(f"  -> Selected: {relative_display(selected)}")
                return selected
        except ValueError:
            pass
        print("  Invalid choice, try again.")


def select_template_and_database() -> tuple[Path, Path]:
    """Scan BASE + subfolders and let the user pick Template + Consolidated file."""
    print("Scanning for Excel files (this may take a moment with many files)...")
    all_xlsx = find_xlsx_files(BASE)

    templates = [p for p in all_xlsx if is_template(p)]
    consolidates = [p for p in all_xlsx if is_consolidated(p)]

    if not templates:
        print("WARNING: No obvious Template files found. Showing all .xlsx files.")
        templates = all_xlsx
    if not consolidates:
        print("WARNING: No obvious Consolidated files found. Showing all .xlsx files.")
        consolidates = all_xlsx

    print(f"Found {len(templates)} possible Template(s) and {len(consolidates)} possible Consolidated file(s).")

    template = choose_from_list("SELECT DELIVERY NOTE TEMPLATE", templates)
    database = choose_from_list("SELECT CONSOLIDATED / DATABASE FILE", consolidates)

    return template, database


# ---------------------------------------------------------------------------
# Reading the item lines for a PO
# ---------------------------------------------------------------------------

def load_items(database: Path, po: str) -> pd.DataFrame:
    """Load non-freight lines for the given PO, sorted by original LINE."""
    df = pd.read_excel(database, header=1)

    required = {"PO NO.", "DESCRIPTION", "LINE", "UOM", "QTY"}
    missing = required - set(df.columns)
    if missing:
        raise RuntimeError(f"Database is missing columns: {sorted(missing)}")

    po = str(po).strip()
    result = df[df["PO NO."].astype(str).str.strip() == po].copy()

    freight_mask = (
        result["DESCRIPTION"]
        .astype(str)
        .str.upper()
        .str.contains(r"\bFREIGHT\b", regex=True, na=False)
    )
    result = result[~freight_mask]

    if result.empty:
        raise RuntimeError(f"No non-freight items found for PO {po}")

    return result.sort_values("LINE").reset_index(drop=True)


# ---------------------------------------------------------------------------
# Writing the delivery note
# ---------------------------------------------------------------------------

def excel_copy_row_format(ws, source_row: int, target_row: int) -> None:
    """Copy formatting (not values) from source_row to target_row."""
    source = ws.Range(
        ws.Cells(source_row, TABLE_FIRST_COL),
        ws.Cells(source_row, TABLE_LAST_COL),
    )
    target = ws.Range(
        ws.Cells(target_row, TABLE_FIRST_COL),
        ws.Cells(target_row, TABLE_LAST_COL),
    )
    source.Copy()
    target.PasteSpecial(-4122)          # xlPasteFormats
    ws.Rows(target_row).RowHeight = ws.Rows(source_row).RowHeight


def fill_po_labels(ws, po: str) -> None:
    """
    Find every label in the header area (rows 1..PO_LABEL_SEARCH_ROWS) that
    contains "PO NUMBER" or "PO#" (e.g. "PO NUMBER:", "ABRAJ AL-DAR PO#")
    and append the PO number to it. Works the same across all 4 templates
    without needing to know each one's exact wording.
    """
    for row in ws.Range(
        ws.Cells(1, 1), ws.Cells(PO_LABEL_SEARCH_ROWS, TABLE_LAST_COL)
    ).Rows:
        for cell in row.Cells:
            value = cell.Value
            if isinstance(value, str) and PO_LABEL_PATTERN.search(value):
                cell.Value = value.rstrip() + " " + po


def make_delivery_note(excel, template: Path, database: Path, po: str) -> Path:
    po = str(po).strip()
    items = load_items(database, po)

    output = BASE / f"Delivery Note {po}.xlsx"
    if output.exists():
        output.unlink()

    wb = excel.Workbooks.Open(
        str(template),
        UpdateLinks=0,
        ReadOnly=False,
        IgnoreReadOnlyRecommended=True,
        AddToMru=False,
    )

    try:
        ws = wb.Worksheets(1)

        fill_po_labels(ws, po)

        n = len(items)
        extra_rows = max(0, n - 1)

        if extra_rows:
            ws.Rows(f"{FOOTER_ROW}:{FOOTER_ROW + extra_rows - 1}").Insert(
                Shift=-4121,          # xlDown
                CopyOrigin=0,         # xlFormatFromLeftOrAbove
            )
            for i in range(1, n):
                excel_copy_row_format(ws, ITEM_ROW, ITEM_ROW + i)

        for i, (_, item) in enumerate(items.iterrows()):
            r = ITEM_ROW + i

            ws.Cells(r, 2).Value = i + 1                # B = SR#
            ws.Cells(r, 3).Value = item["PO NO."]       # C = PO#
            ws.Cells(r, 5).Value = item["DESCRIPTION"]  # E = description
            ws.Cells(r, 7).Value = item["UOM"]          # G = UOM
            ws.Cells(r, 8).Value = item["QTY"]          # H = PO quantity
            ws.Cells(r, 9).Value = item["LINE"]         # I = original PO line

            for col in range(TABLE_FIRST_COL, TABLE_LAST_COL + 1):
                cell = ws.Cells(r, col)
                cell.Font.Name = "Calibri"
                cell.Font.Size = 11
                cell.Font.Bold = False
                cell.HorizontalAlignment = -4108   # xlCenter
                cell.VerticalAlignment = -4108     # xlCenter
                cell.WrapText = True

        item_range = ws.Range(
            ws.Cells(ITEM_ROW, TABLE_FIRST_COL),
            ws.Cells(ITEM_ROW + n - 1, TABLE_LAST_COL),
        )
        item_range.Rows.AutoFit()

        wb.SaveAs(
            str(output),
            FileFormat=51,          # xlOpenXMLWorkbook
            ConflictResolution=2,
        )
        wb.Save()

    finally:
        wb.Close(SaveChanges=False)

    time.sleep(0.5)
    validate_output(template, output, n)

    print(f"SUCCESS: {output}")
    print(f"Items: {n}")
    print("Images preserved.")
    print("Validation passed.")
    return output


# ---------------------------------------------------------------------------
# Post-generation validation
# ---------------------------------------------------------------------------

def validate_output(template: Path, output: Path, n_items: int) -> None:
    """Sanity-check the generated file: valid ZIP, images intact, header
    unchanged, and exactly n_items rows of data written."""

    with zipfile.ZipFile(output, "r") as z:
        bad = z.testzip()
        if bad:
            raise RuntimeError(f"Generated XLSX ZIP is corrupt: {bad}")

    template_check = openpyxl.load_workbook(template, data_only=False)
    output_check = openpyxl.load_workbook(output, data_only=False)

    try:
        tw = template_check.active
        ow = output_check.active

        if len(ow._images) != len(tw._images):
            raise RuntimeError(
                f"IMAGE VALIDATION FAILED: template={len(tw._images)}, output={len(ow._images)}"
            )

        if ow["A1"].value != tw["A1"].value:
            raise RuntimeError("TEMPLATE VALIDATION FAILED: A1 changed.")

        written = 0
        for r in range(ITEM_ROW, ITEM_ROW + n_items + 5):
            if ow.cell(r, 5).value is not None:
                written += 1
            else:
                break
        if written != n_items:
            raise RuntimeError(
                f"DATA VALIDATION FAILED: expected {n_items} item rows, found {written}"
            )

    finally:
        template_check.close()
        output_check.close()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    try:
        template, database = select_template_and_database()
    except RuntimeError as e:
        print(f"ERROR: {e}")
        return

    print()
    print("-" * 70)
    print(f"Template : {relative_display(template)}")
    print(f"Database : {relative_display(database)}")
    print("-" * 70)
    print()
    print("Enter PO numbers. Type START when finished.")
    print()

    pos = []
    while True:
        value = input("> ").strip()
        if value.lower() == "start":
            break
        if value:
            pos.append(value)

    if not pos:
        print("No PO numbers entered.")
        return

    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.ScreenUpdating = False
    excel.EnableEvents = False

    failed = False
    try:
        for po in pos:
            try:
                make_delivery_note(excel, template, database, po)
            except Exception as exc:
                failed = True
                print()
                print(f"ERROR for PO {po}:")
                print(exc)
                print()
    finally:
        excel.Quit()

    if not failed:
        print()
        print("ALL DELIVERY NOTES CREATED AND VALIDATED.")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\nFATAL UNHANDLED ERROR: {e}")
    finally:
        input("\nPress Enter to exit...")