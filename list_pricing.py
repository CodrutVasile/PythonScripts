import sys
import shutil
from datetime import datetime, date
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border
from pathlib import Path


def get_base_dir():
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).parent.resolve()
    else:
        return Path(__file__).parent.resolve()


def copy_cell_formatting(src_cell, dst_cell):
    if src_cell.has_style:
        dst_cell.font = Font(
            name=src_cell.font.name,
            size=src_cell.font.size,
            bold=src_cell.font.bold,
            italic=src_cell.font.italic,
            color=src_cell.font.color,
        )
        dst_cell.border = Border(
            left=src_cell.border.left,
            right=src_cell.border.right,
            top=src_cell.border.top,
            bottom=src_cell.border.bottom,
        )
        dst_cell.fill = PatternFill(
            fill_type=src_cell.fill.fill_type,
            start_color=src_cell.fill.start_color,
            end_color=src_cell.fill.end_color,
        )
        dst_cell.alignment = Alignment(
            horizontal=src_cell.alignment.horizontal,
            vertical=src_cell.alignment.vertical,
            wrap_text=src_cell.alignment.wrap_text,
        )


def is_freight_line(description: str) -> bool:
    if not description:
        return False
    desc = description.strip().lower()
    return desc in ("freight", "freight charges", "freight charge")


def find_file_by_name(filename: str):
    root = get_base_dir()
    for path in root.rglob(filename):
        if not path.name.startswith("~$"):
            return path
    return None


def find_consolidated_files():
    root = get_base_dir()
    files = []
    for path in root.rglob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        if path.name.lower().startswith("consolidated"):
            files.append(path)
    files.sort(key=lambda p: (str(p.parent).lower(), p.name.lower()))
    return files


def select_database_file():
    """
    Show the company menu in a fixed order.

    The consolidated workbook can live anywhere underneath the program folder.
    We identify it by the company name instead of exposing the full filename/path
    to the user.
    """
    companies = [
        ("Abraj", "Consolidated Abraj.xlsx"),
        ("Ameedat", "Consolidated Ameedat.xlsx"),
        ("Athab", "Consolidated Athab.xlsx"),
        ("Wardat", "Consolidated Wardat.xlsx"),
        ("Wasmat Baghdad", "Consolidated Wasmat Baghdad.xlsx"),
    ]

    while True:
        print("\nSelect Company: Type a number then press ENTER or Type exit then press ENTER\n")
        for idx, (company_name, _) in enumerate(companies, start=1):
            print(f"{idx}. {company_name}")

        choice = input("\n").strip()

        if choice.lower() == "exit":
            return "exit"

        try:
            choice_idx = int(choice) - 1
            if not (0 <= choice_idx < len(companies)):
                print("Invalid choice.")
                continue
        except ValueError:
            print("Invalid input.")
            continue

        company_name, filename = companies[choice_idx]
        path = find_file_by_name(filename)

        if path is None:
            print(f"\nERROR: {filename} was not found.")
            print("Make sure the consolidated workbooks are inside the folder/subfolders.")
            continue

        return str(path)


def process_multiple_pos(po_list, baza_file):
    template_list = find_file_by_name("Template List of Items.xlsx")
    template_pricing = find_file_by_name("Template Pricing.xlsx")

    if template_list is None:
        print("\nTemplate List of Items.xlsx not found.")
        return []

    if template_pricing is None:
        print("\nTemplate Pricing.xlsx not found.")
        return []

    try:
        wb_baza = load_workbook(baza_file, data_only=True)
        ws_baza = wb_baza.active
    except Exception as e:
        print(f"\nError reading database: {e}")
        return []

    headers = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in ws_baza[2]
    ]

    try:
        po_col_idx = headers.index("PO NO.")
        desc_col_idx = headers.index("DESCRIPTION")
        item_col_idx = headers.index("LINE")
        qty_col_idx = headers.index("QTY")
        uom_col_idx = headers.index("UOM")
        edd_col_idx = headers.index("EDD")
    except ValueError as e:
        print(f"\nMissing column: {e}")
        return []

    base_dir = get_base_dir()

    # Keep the results in memory so the console can print the requested
    # "LOG" and "GENERATED" sections separately.
    results = []
    generated_list_files = []
    generated_pricing_files = []

    for po_input in po_list:
        po_input = str(po_input).strip()
        safe_po_filename = po_input.replace("/", "-").replace("\\", "-").replace(" ", "")

        po_items = []
        freight_skipped = 0

        for row in ws_baza.iter_rows(min_row=3, values_only=True):
            if len(row) <= max(
                po_col_idx, desc_col_idx, qty_col_idx, uom_col_idx, edd_col_idx
            ):
                continue

            row_po = str(row[po_col_idx]).strip() if row[po_col_idx] is not None else ""
            if row_po != po_input:
                continue

            desc = str(row[desc_col_idx]).strip() if row[desc_col_idx] is not None else ""

            if is_freight_line(desc):
                freight_skipped += 1
                continue

            po_items.append(row)

        results.append({
            "po": po_input,
            "freight_skipped": freight_skipped,
            "item_count": len(po_items),
        })

        # Nothing to generate for this PO.
        if not po_items:
            continue

        output_list_name = base_dir / f"List of items {safe_po_filename}.xlsx"
        output_pricing_name = base_dir / f"Pricing {safe_po_filename}.xlsx"

        # ------------------------------------------------------------
        # LIST OF ITEMS
        # ------------------------------------------------------------
        shutil.copy2(template_list, output_list_name)
        wb_list = load_workbook(output_list_name)
        ws_list = wb_list.active

        for idx, row in enumerate(po_items, start=1):
            excel_row = idx + 1

            ws_list[f"A{excel_row}"] = idx
            ws_list[f"B{excel_row}"] = row[desc_col_idx]
            ws_list[f"C{excel_row}"] = row[qty_col_idx]
            ws_list[f"D{excel_row}"] = row[uom_col_idx]

            edd_val = row[edd_col_idx]

            if isinstance(edd_val, str):
                edd_val_str = edd_val.strip()
                parsed_date = None

                for fmt in [
                    "%Y-%m-%d %H:%M:%S",
                    "%Y-%m-%d",
                    "%m/%d/%y",
                    "%m/%d/%Y",
                    "%d/%m/%Y",
                ]:
                    try:
                        parsed_date = datetime.strptime(edd_val_str, fmt)
                        break
                    except ValueError:
                        continue

                if parsed_date:
                    edd_val = parsed_date

            ws_list[f"E{excel_row}"] = edd_val

            if isinstance(edd_val, (datetime, date)):
                ws_list[f"E{excel_row}"].number_format = "mm/dd/yy"

            if excel_row > 2:
                for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I"]:
                    copy_cell_formatting(
                        ws_list[f"{col_letter}2"],
                        ws_list[f"{col_letter}{excel_row}"]
                    )

        wb_list.save(output_list_name)
        wb_list.close()

        # ------------------------------------------------------------
        # PRICING
        # ------------------------------------------------------------
        shutil.copy2(template_pricing, output_pricing_name)
        wb_pricing = load_workbook(output_pricing_name)

        ws_price = (
            wb_pricing["Pricing Quote"]
            if "Pricing Quote" in wb_pricing.sheetnames
            else wb_pricing.active
        )

        START_ROW = 13

        if len(po_items) > 1:
            ws_price.insert_rows(START_ROW + 1, amount=len(po_items) - 1)

        template_row = START_ROW

        for item_index, row in enumerate(po_items):
            current_row = START_ROW + item_index
            sequential_no = item_index + 1

            if current_row != template_row:
                for col in range(1, ws_price.max_column + 1):
                    src_cell = ws_price.cell(row=template_row, column=col)
                    dst_cell = ws_price.cell(row=current_row, column=col)

                    if src_cell.has_style:
                        dst_cell._style = copy(src_cell._style)

                    dst_cell.font = copy(src_cell.font)
                    dst_cell.fill = copy(src_cell.fill)
                    dst_cell.border = copy(src_cell.border)
                    dst_cell.alignment = copy(src_cell.alignment)
                    dst_cell.number_format = src_cell.number_format
                    dst_cell.protection = copy(src_cell.protection)

                    if src_cell.data_type == "f" and src_cell.value:
                        formula = src_cell.value
                        formula = formula.replace(str(template_row), str(current_row))
                        dst_cell.value = formula

                if template_row in ws_price.row_dimensions:
                    ws_price.row_dimensions[current_row].height = (
                        ws_price.row_dimensions[template_row].height
                    )

            ws_price[f"A{current_row}"] = sequential_no
            ws_price[f"E{current_row}"] = row[desc_col_idx]
            ws_price[f"F{current_row}"] = row[qty_col_idx]
            ws_price[f"G{current_row}"] = row[uom_col_idx]

        wb_pricing.save(output_pricing_name)
        wb_pricing.close()

        generated_list_files.append(output_list_name.name)
        generated_pricing_files.append(output_pricing_name.name)

    # ------------------------------------------------------------
    # CONSOLE LOG
    # ------------------------------------------------------------
    print("\n---")
    print("\nLOG")

    for result in results:
        print(f"\n{result['po']}")

        if result["freight_skipped"]:
            print(f"  Skipped {result['freight_skipped']} freight line(s)")

        if result["item_count"]:
            print(f"  Found {result['item_count']} items")
        else:
            print("  No items found")

    print("\nLOG")
    print("\n---")

    for filename in generated_list_files:
        print(f"  → GENERATED {filename}")

    for filename in generated_pricing_files:
        print(f"  → GENERATED {filename}")

    return generated_list_files + generated_pricing_files


if __name__ == "__main__":
    while True:
        baza_file = select_database_file()

        if baza_file == "exit":
            print("\nExiting.")
            break

        if baza_file is None:
            input("\nPress Enter to continue...")
            continue

        print("\nEnter one PO number per line.")
        print("Type a number then press ENTER")
        print("When all PO's are written, type start on a new line then press ENTER\n")

        input_pos = []

        while True:
            line = input().strip()

            if line.lower() == "exit":
                print("\nExiting.")
                sys.exit(0)

            if line.lower() == "start":
                break

            if line:
                # Normally one PO per line. If several numeric PO values are
                # accidentally pasted on one line, accept them too.
                parts = line.split()
                if len(parts) > 1 and all(part.isdigit() for part in parts):
                    input_pos.extend(parts)
                else:
                    input_pos.append(line)

        if not input_pos:
            print("\nNo PO numbers entered.")
            continue

        try:
            process_multiple_pos(input_pos, baza_file)
        except Exception as e:
            print(f"\nError: {e}")
