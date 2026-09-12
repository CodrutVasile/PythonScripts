#!/usr/bin/env python3
import sys
import re
from pathlib import Path
import openpyxl

SOURCE_FOLDER_NAME = "Extract RFQs in Consolidated"
CONSOLIDATED_FOLDER_HINT = "consolidat"  

HEADER_LABELS = {
    'item no.': 'line', 
    'qty': 'quantity', 
    'ui': 'uom', 
    'description': 'description', 
    'long description': 'long_description',
    'unit price': 'unit_cost', 
    'delivery time': 'delivery_date', 
    'totals': 'totals', 
    'part #': 'part_number'
}

def normalize(text):
    if text is None:
        return ''
    return re.sub(r'\s+', ' ', str(text)).strip().lower()

def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)

def find_header_columns(ws):
    found = {}
    for row in ws.iter_rows():
        for cell in row:
            label = normalize(cell.value)
            if label in HEADER_LABELS and HEADER_LABELS[label] not in found:
                found[HEADER_LABELS[label]] = cell.column
        if len(found) == len(set(HEADER_LABELS.values())):
            return found
    return found

def is_header_row(ws, row_idx, cols):
    for col in cols.values():
        val = normalize(ws.cell(row=row_idx, column=col).value)
        if val in HEADER_LABELS:
            return True
    return False

def extract_records(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    all_records = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        cols = find_header_columns(ws)
        required = {'uom', 'description', 'quantity', 'unit_cost'}
        if not required.issubset(cols):
            continue

        desc_col = cols['description']
        qty_col = cols['quantity']
        uom_col = cols['uom']
        cost_col = cols['unit_cost']
        date_col = cols.get('delivery_date')
        line_col = cols.get('line')
        long_desc_col = cols.get('long_description')
        part_col = cols.get('part_number')
        totals_col = cols.get('totals')
        header_row = min(cols.values())
        start_row = header_row + 1
        max_row = ws.max_row
        for r in range(start_row, max_row + 1):
            if is_header_row(ws, r, cols):
                continue

            qty_val = ws.cell(row=r, column=qty_col).value
            uom_val = ws.cell(row=r, column=uom_col).value
            cost_val = ws.cell(row=r, column=cost_col).value
            desc_val = ws.cell(row=r, column=desc_col).value
            if not (is_number(qty_val) and uom_val and (desc_val not in (None, ''))):
                continue

            if not is_number(cost_val) and totals_col:
                totals_val = ws.cell(row=r, column=totals_col).value
                if is_number(totals_val) and qty_val:
                    cost_val = totals_val / qty_val

            if is_number(cost_val):
                desc_parts = []
                if part_col:
                    part_val = ws.cell(row=r, column=part_col).value
                    if part_val not in (None, ''):
                        desc_parts.append(f'PART# -{part_val}')
                desc_parts.append(str(desc_val))
                if long_desc_col:
                    long_val = ws.cell(row=r, column=long_desc_col).value
                    if long_val not in (None, ''):
                        desc_parts.append(str(long_val))
                desc_clean = re.sub(r'\s+', ' ', ' '.join(desc_parts)).strip()
                line_val = ws.cell(row=r, column=line_col).value if line_col else None
                date_val = ws.cell(row=r, column=date_col).value if date_col else None
                all_records.append({
                    'Line': line_val,
                    'Part / Description': desc_clean,
                    'Quantity': qty_val,
                    'UOM': str(uom_val).strip(),
                    'Unit Cost': cost_val,
                    'Delivery Date': date_val
                })
    return all_records

def extract_rfq_number(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    raw_value = None
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows():
            for cell in row:
                if normalize(cell.value) == 'rfq control no.':
                    col = cell.column
                    r = cell.row
                    for cc in range(col + 1, col + 8):
                        v = ws.cell(row=r, column=cc).value
                        if v not in (None, ''):
                            raw_value = str(v).strip()
                            break
                    if raw_value:
                        continue
                    for rr in range(r + 1, r + 4):
                        v = ws.cell(row=rr, column=col).value
                        if v not in (None, ''):
                            raw_value = str(v).strip()
                            break
            if raw_value:
                break
        if raw_value:
            break

    if not raw_value:
        m = re.search(r'(\d{4,})', path.stem)
        if m:
            raw_value = m.group(1)

    if not raw_value:
        return None

    digits = re.findall(r'\d+', raw_value)
    if digits:
        return int(digits[0])
    return raw_value

DB_COLS = {'line': 'N', 'description': 'O', 'quantity': 'P', 'uom': 'Q', 'unit_cost': 'R', 'delivery_date': 'S', 'po_number': 'L', 'rfq_number': 'E'}

def find_last_data_row(ws, anchor_col_letter='N'):
    for r in range(ws.max_row, 0, -1):
        if ws[f'{anchor_col_letter}{r}'].value not in (None, ''):
            return r
    return 1

def append_entries_to_database(db_path, entries):
    wb = openpyxl.load_workbook(db_path, data_only=False)
    ws = wb.active
    last_row = find_last_data_row(ws, DB_COLS['line'])
    for entry in entries:
        new_row = last_row + 1
        ws[f"{DB_COLS['line']}{new_row}"] = entry['line']
        ws[f"{DB_COLS['description']}{new_row}"] = entry['description']
        ws[f"{DB_COLS['quantity']}{new_row}"] = entry['quantity']
        ws[f"{DB_COLS['uom']}{new_row}"] = entry['uom']
        ws[f"{DB_COLS['unit_cost']}{new_row}"] = entry['unit_cost']
        ws[f"{DB_COLS['delivery_date']}{new_row}"] = entry['delivery_date']
        ws[f"{DB_COLS['rfq_number']}{new_row}"] = entry['rfq_number']
        last_row = new_row
    wb.save(db_path)
    return last_row

# ---------------------------------------------------------------------------
# Locating folders & interactive prompt UI matching extract_complex.py
# ---------------------------------------------------------------------------

def find_named_folder(start_dir, folder_name):
    start_dir = start_dir.resolve()
    target = folder_name.lower()

    direct = start_dir / folder_name
    if direct.is_dir():
        return direct
    for child in start_dir.iterdir():
        if child.is_dir() and child.name.lower() == target:
            return child

    for p in start_dir.rglob('*'):
        if p.is_dir() and p.name.lower() == target:
            return p

    for parent in start_dir.parents:
        for child in parent.iterdir():
            if child.is_dir() and child.name.lower() == target:
                return child
        for p in parent.rglob('*'):
            if p.is_dir() and p.name.lower() == target:
                return p

    return None

def find_consolidated_folder(start_dir, source_folder):
    search_roots = [start_dir.resolve()]
    if source_folder is not None:
        search_roots.append(source_folder.parent)

    for root in search_roots:
        try:
            children = list(root.iterdir())
        except OSError:
            continue
        for child in children:
            if child.is_dir() and CONSOLIDATED_FOLDER_HINT in child.name.lower():
                return child
    return None

def find_database_files(consolidated_folder):
    return sorted(
        p for p in consolidated_folder.rglob('*.xlsx')
        if not p.name.startswith('~$') and CONSOLIDATED_FOLDER_HINT in p.stem.lower()
    )

def find_source_files(folder):
    return sorted(p for p in folder.glob('*.xlsx') if not p.name.startswith('~$'))

def prompt_select_database(db_files):
    """Presents a numbered menu matching extract_complex.py layout format."""
    print("Select Company (Type the corresponding number PRESS ENTER PRESS START PRESS ENTER)\n")
    print("Select target database for ALL files:")
    for idx, db_path in enumerate(db_files, start=1):
        print(f"  {idx}. {db_path.name}")

    while True:
        choice = input(f"Enter choice (1-{len(db_files)}): ").strip()
        if choice.isdigit():
            idx = int(choice)
            if 1 <= idx <= len(db_files):
                return db_files[idx - 1]
        print(f"Invalid selection. Enter a number between 1 and {len(db_files)}.")

def run_once(script_dir):
    source_folder = find_named_folder(script_dir, SOURCE_FOLDER_NAME)
    if source_folder is None:
        print(f"\nCould not find a folder named '{SOURCE_FOLDER_NAME}'.")
        print("Make sure it exists near this script.")
        return

    source_files = find_source_files(source_folder)
    if not source_files:
        print(f"No .xlsx files found in '{source_folder}'.")
        return

    consolidated_folder = find_consolidated_folder(script_dir, source_folder)
    if consolidated_folder is None:
        print("\nCould not find the Consolidated folder.")
        return

    db_files = find_database_files(consolidated_folder)
    if not db_files:
        print(f"No Consolidated_*.xlsx database files found in '{consolidated_folder}'.")
        return

    target_db = prompt_select_database(db_files)
    pending_records = []

    print(f"\nProcessing {len(source_files)} source file(s) in '{source_folder.name}':")
    for f in source_files:
        rfq_number = extract_rfq_number(f)
        records = extract_records(f)
        if not records:
            print(f"  {f.name}: no line items found, skipping")
            continue

        for rec in records:
            pending_records.append({
                'line': rec['Line'],
                'description': rec['Part / Description'],
                'quantity': rec['Quantity'],
                'uom': rec['UOM'],
                'unit_cost': rec['Unit Cost'],
                'delivery_date': rec['Delivery Date'],
                'rfq_number': rfq_number
            })
        print(f"  {f.name}: RFQ #{rfq_number} -> {target_db.name}, {len(records)} line item(s)")

    if not pending_records:
        print("\nNo line items to add.")
        return

    print()
    last_row = append_entries_to_database(target_db, pending_records)
    print(f"Appended {len(pending_records)} row(s) from {len(source_files)} file(s) to '{target_db.name}' (now ends at row {last_row}).")

def get_app_dir():
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent

def main():
    script_dir = get_app_dir()
    while True:
        run_once(script_dir)
        choice = input("Type EXIT  then PRESS  enter to quit, \n").strip().lower()
        if choice == "exit":
            print("Exitting")
            break

if __name__ == '__main__':
    main()